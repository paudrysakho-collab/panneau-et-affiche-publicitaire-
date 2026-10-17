"""Génère le PPTX modifiable d'une version de l'affiche (sombre par défaut, ou --clair).

Fond = rendu Remotion de la composition « PanneauFond » / « PanneauClairFond » (sans texte ni logo),
puis le logo en image déplaçable et chaque bloc de src/textes.json en zone de texte modifiable.

Usage (depuis panneau/) :
    npx remotion still src/index.ts PanneauFond exports/fond.png --scale=2.829
    python3 tools/make_pptx.py
    npx remotion still src/index.ts PanneauClairFond exports/fond-clair.png --scale=2.829
    python3 tools/make_pptx.py --clair
Variante de couleur (nom de palette : safran, corail, anis, turquoise, framboise…) :
    npx remotion still src/index.ts PanneauOcreFond exports/fond-safran.png --scale=2.829
    python3 tools/make_pptx.py --palette safran
"""
import sys
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Mm, Pt

RACINE = Path(__file__).resolve().parent.parent
LARGEUR_PX, HAUTEUR_PX = 1240, 1754
LARGEUR_MM = 594.0
MM_PAR_PX = LARGEUR_MM / LARGEUR_PX
PT_PAR_PX = MM_PAR_PX / 0.352778

# Mêmes valeurs que C et CLAIR dans src/Panneau.tsx (cremeDoux aplati sur le fond).
COULEURS = {
    'or': (0xCF, 0xA5, 0x48),
    'orPale': (0xE6, 0xCF, 0x98),
    'creme': (0xF4, 0xEC, 0xE0),
    'cremeDoux': (0xC8, 0xC1, 0xB7),
}
COULEURS_CLAIR = {
    'or': (0xA7, 0x7A, 0x26),
    'orPale': (0x7D, 0x15, 0x30),
    'creme': (0x2A, 0x16, 0x19),
    'cremeDoux': (0x5E, 0x4E, 0x4D),
}

# Variantes de couleur : mêmes valeurs que PALETTES dans src/Panneau.tsx.
COULEURS_VARIANTES = {
    'grenat': {
        'or': (0xEA, 0xC4, 0x70),
        'orPale': (0xF6, 0xE3, 0xB0),
        'creme': (0xFF, 0xF7, 0xEC),
        'cremeDoux': (0xC8, 0xBD, 0xB6),
    },
    'vert': {
        'or': (0xDC, 0xB7, 0x5E),
        'orPale': (0xEF, 0xE0, 0xAA),
        'creme': (0xF5, 0xF2, 0xE6),
        'cremeDoux': (0xB6, 0xB7, 0xAD),
    },
    'bleu': {
        'or': (0xE2, 0xBC, 0x63),
        'orPale': (0xF1, 0xDF, 0xA9),
        'creme': (0xF1, 0xF3, 0xEF),
        'cremeDoux': (0xB3, 0xB6, 0xB6),
    },
    'safran': {
        'or': (0x8C, 0x22, 0x13),
        'orPale': (0x9C, 0x3A, 0x1C),
        'creme': (0x2A, 0x11, 0x08),
        'cremeDoux': (0x5D, 0x3C, 0x16),
    },
    'corail': {
        'or': (0x73, 0x16, 0x2A),
        'orPale': (0x8C, 0x2A, 0x36),
        'creme': (0x3B, 0x0F, 0x10),
        'cremeDoux': (0x6A, 0x2D, 0x27),
    },
    'anis': {
        'or': (0x45, 0x61, 0x0F),
        'orPale': (0x55, 0x73, 0x1A),
        'creme': (0x1E, 0x2B, 0x06),
        'cremeDoux': (0x49, 0x58, 0x19),
    },
    'turquoise': {
        'or': (0x0D, 0x55, 0x5E),
        'orPale': (0x13, 0x68, 0x72),
        'creme': (0x07, 0x26, 0x2A),
        'cremeDoux': (0x1E, 0x50, 0x55),
    },
    'framboise': {
        'or': (0x7D, 0x17, 0x38),
        'orPale': (0x9A, 0x24, 0x48),
        'creme': (0x3A, 0x0A, 0x1C),
        'cremeDoux': (0x69, 0x28, 0x3E),
    },
}

POLICES = {'serif': 'Cormorant Garamond', 'sans': 'Montserrat'}
ALIGN = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}


def px(v):
    return Mm(v * MM_PAR_PX)


def main():
    # --clair pour la version crème, --palette <nom> pour une variante de couleur
    clair = '--clair' in sys.argv
    variante = sys.argv[sys.argv.index('--palette') + 1] if '--palette' in sys.argv else None
    if variante:
        couleurs = COULEURS_VARIANTES[variante]
        suffixe = f'-{variante}'
    else:
        couleurs = COULEURS_CLAIR if clair else COULEURS
        suffixe = '-clair' if clair else ''
    textes = json.loads((RACINE / 'src/textes.json').read_text(encoding='utf-8'))

    fond_png = RACINE / f'exports/fond{suffixe}.png'
    fond_jpg = RACINE / f'exports/fond{suffixe}.jpg'
    Image.open(fond_png).convert('RGB').save(fond_jpg, quality=92)

    prs = Presentation()
    prs.slide_width = px(LARGEUR_PX)
    prs.slide_height = px(HAUTEUR_PX)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    slide.shapes.add_picture(str(fond_jpg), 0, 0, prs.slide_width, prs.slide_height)
    slide.shapes.add_picture(str(RACINE / 'public/logo_cave71_bloc.png'), px(240), px(98), width=px(760))

    for bloc in textes['blocs']:
        interligne = bloc.get('interligne', 1.2)
        hauteur = sum(max(s['taille'] for s in ligne) * interligne for ligne in bloc['lignes'])
        boite = slide.shapes.add_textbox(px(bloc['x']), px(bloc['y']), px(bloc['w']), px(hauteur))
        boite.name = bloc['id']
        tf = boite.text_frame
        tf.word_wrap = False
        tf.auto_size = None
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)

        for i, ligne in enumerate(bloc['lignes']):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = ALIGN[bloc['align']]
            para.line_spacing = Pt(max(seg['taille'] for seg in ligne) * interligne * PT_PAR_PX)
            for seg in ligne:
                run = para.add_run()
                run.text = seg['t']
                f = run.font
                f.name = POLICES[seg['police']]
                f.size = Pt(seg['taille'] * PT_PAR_PX)
                f.bold = seg.get('graisse', 400) >= 600
                f.italic = bool(seg.get('italique'))
                f.color.rgb = RGBColor(*couleurs[seg['couleur']])
                if seg.get('espacement'):
                    run._r.get_or_add_rPr().set('spc', str(round(seg['espacement'] * PT_PAR_PX * 100)))
                rpr = run._r.get_or_add_rPr()
                for tag in ('a:latin', 'a:cs'):
                    el = rpr.find(qn(tag))
                    if el is None:
                        el = rpr.makeelement(qn(tag), {})
                        rpr.append(el)
                    el.set('typeface', POLICES[seg['police']])

    sortie = (
        RACINE / f'exports/variantes/cave71-A1-{variante}-modifiable.pptx'
        if variante
        else RACINE / f'exports/panneau-cave71-A1{suffixe}-modifiable.pptx'
    )
    prs.save(sortie)
    print(sortie)


if __name__ == '__main__':
    main()
